# Required libraries: pip install python-pptx requests

import sys
import json
import pptx
import requests
from io import BytesIO
from pptx.util import Inches, Pt
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_DATA_LABEL_POSITION
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_VERTICAL_ANCHOR, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.chart.data import CategoryChartData
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
import textwrap


# --- THEME AND STYLING CONSTANTS ---
THEME_COLORS = {
    'text': RGBColor(0x36, 0x36, 0x36),
    'accent_orange': RGBColor(0xF3, 0x7F, 0x33),
    'accent_teal': RGBColor(0x00, 0x80, 0x80),
    'accent_gray': RGBColor(0x4F, 0x55, 0x5D),
    'chart_green': RGBColor(8, 174, 80),
    'chart_yellow': RGBColor(255, 215, 53),
    'chart_red': RGBColor(253, 9, 7),
}


# --- IMAGE HELPER ---
def _image_stream(image_path):
    try:
        if isinstance(image_path, str) and image_path.startswith(("http://", "https://")):
            resp = requests.get(image_path, timeout=10)
            resp.raise_for_status()
            return BytesIO(resp.content)  # always fresh buffer
        with open(image_path, "rb") as f:
            return BytesIO(f.read())  # read into new buffer
    except Exception as e:
        print(f"Warning: Unable to load image from {image_path} ({e})", file=sys.stderr)
        return None


# --- FOOTER HELPER ---
def _add_footer(prs, slide, slide_index, total_slides, logo_path=None, hide_page=False, is_title_slide=False):
    footer_y = prs.slide_height - Inches(0.6)

    # --- LOGO (skip on title slide) ---
    if logo_path and not is_title_slide:
        img_stream = _image_stream(logo_path)
        if img_stream is not None:
            slide.shapes.add_picture(img_stream, Inches(0.5), footer_y, Inches(1.59), Inches(0.18))

    # --- Page number (conditional) ---
    if not hide_page and slide_index != 0 and slide_index != total_slides - 1:
        width = Inches(0.28) 
        height = Inches(0.25)
        right_margin = Inches(0.38)
        left = prs.slide_width - right_margin - width
        top = footer_y

        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0, 0, 0)
        shape.line.fill.background() 

        text_frame = shape.text_frame
        text_frame.clear()
        text_frame.autofit = False

        p = text_frame.paragraphs[0]
        p.text = str(slide_index + 1)
        p.font.size = Pt(5)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        p.alignment = PP_ALIGN.CENTER
        text_frame.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE


# --- SLIDE HELPERS ---
def _add_standard_header_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    p = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(9), Inches(0.5)).text_frame.paragraphs[0]
    p.text, p.font.size, p.font.bold, p.font.color.rgb = title, Pt(24), True, THEME_COLORS['accent_gray']
    if subtitle:
        p = slide.shapes.add_textbox(Inches(0.5), Inches(0.6), Inches(9), Inches(0.4)).text_frame.paragraphs[0]
        p.text, p.font.size, p.font.color.rgb = subtitle, Pt(14), THEME_COLORS['text']
    return slide


def _create_title_slide(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    primary_logo_path = data.get('logoImagePath') or data.get('siteLogo')
    if primary_logo_path:
        if img_stream := _image_stream(primary_logo_path):
            slide.shapes.add_picture(img_stream, Inches(0.5), Inches(0.6), Inches(1.70), Inches(0.66))

    # --- Site Logo (Top Right) ---
    if site_logo_path := data.get('siteLogo'):
        if img_stream := _image_stream(site_logo_path):
            slide.shapes.add_picture(img_stream, Inches(8), Inches(0.5), Inches(1.82), Inches(0.20))

    # --- Background and Text (unchanged) ---
    if bg_path := data.get('backgroundImagePath'):
        if img_stream := _image_stream(bg_path):
            slide.shapes.add_picture(img_stream, Inches(0), Inches(2.63), Inches(10), Inches(3))
        # if img_stream is not None: #for complete background
        #     pic = slide.shapes.add_picture(img_stream, Inches(0), Inches(0), Inches(10), Inches(5.625))
        #     slide.shapes._spTree.remove(pic._element)
        #     slide.shapes._spTree.insert(2, pic._element)
            
    p = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(8), Inches(0.5)).text_frame.paragraphs[0]
    p.text, p.font.size, p.font.bold = data.get('clientName', ''), Pt(32), True
    p.font.color.rgb = THEME_COLORS['text']
    p = slide.shapes.add_textbox(Inches(0.5), Inches(1.7), Inches(8), Inches(0.5)).text_frame.paragraphs[0]
    p.text, p.font.size, p.font.color.rgb = data.get('reportTitle', ''), Pt(32), THEME_COLORS['text']
    return slide

def _create_section_break_slide(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(26, 41, 50)
    p = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(4), Inches(0.5)).text_frame.paragraphs[0]
    p.text, p.font.size, p.font.bold = str(data.get('year', '')), Pt(56), True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p = slide.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(4), Inches(1)).text_frame.paragraphs[0]
    p.text, p.font.size = data.get('title', ''), Pt(56)
    p.font.color.rgb = RGBColor(255, 255, 255)

    # --- SECTION BREAK IMAGE ---
    section_img = data.get('imagePath')
    if section_img:
        img_stream = _image_stream(section_img)
        if img_stream is not None:
            slide.shapes.add_picture(img_stream, Inches(6.8), Inches(0.5), Inches(3.2), Inches(4.7))

    return slide


def _hex_to_rgb(hex_color: str):
    hex_color = hex_color.lstrip('#')
    return tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))


def set_cell_border(cell, hex_color="#000000", border_width=None, border_type="solid", sides=("top", "bottom", "left", "right")):
    rgb = _hex_to_rgb(hex_color)  # (r,g,b)
    hex_str = "%02X%02X%02X" % rgb
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()

    side_map = {"left": "a:lnL", "right": "a:lnR", "top": "a:lnT", "bottom": "a:lnB"}
    for side in sides:
        line_dir = side_map.get(side)
        if not line_dir:
            continue

        ln = tcPr.find(qn(line_dir))
        if ln is None:
            ln = OxmlElement(line_dir)
            tcPr.append(ln)

        # Clear previous children
        for child in list(ln):
            ln.remove(child)

        # Width (defaults)
        if border_width is None:
            w = "12700" if border_type == "solid" else "6350"
        else:
            w = str(border_width)
        ln.set("w", w)

        # Dotted vs solid
        if border_type == "dotted":
            prstDash = OxmlElement("a:prstDash")
            prstDash.set("val", "sysDot")
            ln.append(prstDash)

        # Color
        solidFill = OxmlElement("a:solidFill")
        srgbClr = OxmlElement("a:srgbClr")
        srgbClr.set("val", hex_str)
        solidFill.append(srgbClr)
        ln.append(solidFill)


def style_table_borders(table, outer_color="#000000", inner_color="#000000"):
    rows = len(table.rows)
    cols = len(table.columns)

    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            if r == 0:
                set_cell_border(cell, outer_color, border_type="solid", sides=("top",))
            if r == rows - 1:
                set_cell_border(cell, outer_color, border_type="solid", sides=("bottom",))
            if c == 0:
                set_cell_border(cell, outer_color, border_type="solid", sides=("left",))
            if c == cols - 1:
                set_cell_border(cell, outer_color, border_type="solid", sides=("right",))

            # Inner dotted separators
            if r < rows - 1:
                set_cell_border(cell, inner_color, border_type="dotted", sides=("bottom",))
            if c < cols - 1:
                set_cell_border(cell, inner_color, border_type="dotted", sides=("right",))


def _create_table_slide(prs, data):
    MAX_ROWS_PER_SLIDE = 11  # adjust based on font/slide size

    table_data = data.get('table', {})
    headers = table_data.get('headers', [])
    rows_data = table_data.get('rows', [])
    colors = table_data.get('colors', {
        "header": "1E353A",
        "row": "F1FCF8",
        "alt_row": "F1FCF8",
        "section": "0F7C7C",
        "text": "000000"
    })

    slides = []
    if headers and rows_data:
        # Split rows into chunks
        for start in range(0, len(rows_data), MAX_ROWS_PER_SLIDE):
            chunk_rows = rows_data[start:start + MAX_ROWS_PER_SLIDE]

            # --- New slide ---
            slide = _add_standard_header_slide(prs, data.get('title'), data.get('subtitle'))
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor(241, 252, 248)  # background
            slides.append(slide)

            rows, cols = len(chunk_rows) + 1, len(headers)
            total_table_width = Inches(9.0)  # Total width for the table
            table_shape = slide.shapes.add_table(
                rows, cols, Inches(0.5), Inches(1.05), Inches(9.0), Inches(0.35 * rows)
            )
            table = table_shape.table

            # --- Dynamic first column width based on header length ---
            char_width = 0.12  # rough width per character at font size ~8
            longest_header = max(len(h) for h in headers)
            first_col_width = Inches(longest_header * char_width)
            table.columns[0].width = first_col_width

            # --- Calculate and set width for remaining columns ---
            if cols > 1:
                remaining_width = total_table_width - first_col_width
                other_col_width = remaining_width // (cols - 1)
                for i in range(1, cols):
                    table.columns[i].width = other_col_width

            # Rounded rectangle background for header (behind table)
            header_bg = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.08), Inches(9.0), Inches(0.35)
            )
            header_bg.fill.solid()
            header_bg.fill.fore_color.rgb = RGBColor.from_string(colors["header"])
            header_bg.line.fill.background()
            header_bg.adjustments[0] = 0.6  # roundness of corners
            slide.shapes._spTree.remove(header_bg._element)
            slide.shapes._spTree.insert(0, header_bg._element)

            # --- Style header row ---
            for i, header in enumerate(headers):
                cell = table.cell(0, i)
                cell.text = header
                cell.fill.background()
                # Fit text properly
                cell.text_frame.word_wrap = True
                cell.text_frame.auto_size = None
                cell.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.text_frame.margin_top = Pt(0)
                cell.text_frame.margin_bottom = Pt(0)
                cell.text_frame.margin_left = Pt(2)
                cell.text_frame.margin_right = Pt(2)

                p = cell.text_frame.paragraphs[0]
                p.font.bold = True
                p.font.size = Pt(8)
                p.font.color.rgb = RGBColor(255, 255, 255)
                p.alignment = PP_ALIGN.CENTER

            # --- Fill table rows ---
            for r_idx, row_data in enumerate(chunk_rows):
                if isinstance(row_data, dict):
                    row_values = row_data.get("values", [])
                    row_color = row_data.get("color")
                else:
                    row_values = row_data
                    row_color = None

                # BUG FIX: This line is now safe and will not crash if row_values is empty.
                is_section_row = (
                    row_values and isinstance(row_values[0], str) and row_values[0].lower().startswith("wellness credit")
                )

                if row_color:
                    rgb_tuple = _hex_to_rgb(row_color)
                    bg_color = RGBColor(*rgb_tuple)
                else:
                    bg_color = (
                        RGBColor.from_string(colors["row"])
                        if r_idx % 2 == 0 else RGBColor.from_string(colors["alt_row"])
                    )

                for c_idx, cell_data in enumerate(row_values):
                    cell = table.cell(r_idx + 1, c_idx)
                    cell.text = str(cell_data)
                    # Fit text properly
                    cell.text_frame.word_wrap = True
                    cell.text_frame.auto_size = None
                    cell.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
                    cell.text_frame.margin_top = Pt(0)
                    cell.text_frame.margin_bottom = Pt(0)
                    cell.text_frame.margin_left = Pt(2)
                    cell.text_frame.margin_right = Pt(2)

                    p = cell.text_frame.paragraphs[0]
                    p.space_before = Pt(0)
                    p.space_after = Pt(0)
                    p.line_spacing = 1.0
                    p.font.size = Pt(8)

                    if is_section_row:
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = RGBColor.from_string(colors["section"])
                        p.font.bold = True
                        p.font.color.rgb = RGBColor(255, 255, 255)
                        p.alignment = PP_ALIGN.LEFT if c_idx == 0 else PP_ALIGN.CENTER
                    else:
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = bg_color
                        p.font.bold = False
                        p.font.color.rgb = RGBColor.from_string(colors["text"])
                        p.alignment = PP_ALIGN.CENTER if c_idx > 0 else PP_ALIGN.LEFT

            # --- Force consistent row height ---
            for row in table.rows:
                row.height = Inches(0.3)

            style_table_borders(table, outer_color="#000000", inner_color="#000000")

            # --- Footer ---
            if data.get('footer') and start == 0:
                p = slide.shapes.add_textbox(
                    Inches(0.5), Inches(5.2), Inches(9.0), Inches(0.3)
                ).text_frame.paragraphs[0]
                p.text = data.get('footer')
                p.font.size = Pt(6)
                p.font.color.rgb = RGBColor(128, 128, 128)
            #--- Graph Details ---
            text_path = data.get('graphText')
            if text_path:
                txBox = slide.shapes.add_textbox(Inches(0.9), Inches(4.7), Inches(3), Inches(0.3))
                p = txBox.text_frame.paragraphs[0]
                p.text = text_path
                p.font.size = Pt(8)
                p.font.color.rgb = THEME_COLORS['text']

                el = txBox._element
                slide.shapes._spTree.remove(el)
                slide.shapes._spTree.append(el)
            # --- Graph Icon ---
            icon_path = data.get('graphLogo')
            if icon_path:
                img_stream = _image_stream(icon_path)
                if img_stream:
                    pic = slide.shapes.add_picture(img_stream, Inches(0.7), Inches(4.7), width=Inches(0.25))
                    el = pic._element
                    slide.shapes._spTree.remove(el)
                    slide.shapes._spTree.append(el)
            #--- Graph Details ---
            
    return slides

def _create_stacked_bar_chart_slide(prs, data):
    try:
        slide = _add_standard_header_slide(prs, data.get("title", "HRA Results"), "")
        chart_data = CategoryChartData()
        categories = data.get("categories", [])
        chart_data.categories = categories

        def clean_data(series_data):
            cleaned = []
            for value in series_data:
                try:
                    cleaned.append(float(str(value).strip().replace('%', '')) / 100.0)
                except (ValueError, TypeError):
                    cleaned.append(0.0)
            return cleaned

        chart_data.add_series("Low Risk", clean_data(data.get("low", [])))
        chart_data.add_series("Moderate Risk", clean_data(data.get("moderate", [])))
        chart_data.add_series("High Risk", clean_data(data.get("high", [])))
        chart_data.add_series("Very High Risk", clean_data(data.get("very high", [])))

        slide_width = prs.slide_width
        slide_height = prs.slide_height

        longest_label = max((len(str(c)) for c in categories), default=10)
        est_label_width = Inches(0.12 * longest_label)  # ~0.12 inch per char at 8pt

        left = Inches(0.3) + est_label_width
        top = Inches(0.8)
        width = slide_width - left - Inches(0.5)  

        # Dynamic height based on number of categories
        num_cats = len(categories)
        base_height = Inches(4.6)
        min_height = Inches(2.5)
        max_height = slide_height - Inches(2.0)  # leave space for header/footer

        if num_cats > 0:
            height = base_height * (num_cats / 10)  # scale relative to 10 categories
        else:
            height = base_height

        height = max(min_height, min(max_height, height))

        chart_type = XL_CHART_TYPE.BAR_STACKED_100
        chart_shape = slide.shapes.add_chart(chart_type, left, top, width, height, chart_data)
        chart = chart_shape.chart

        # Legend styling (global only, no per-entry support in python-pptx)
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.font.size = Pt(6)
    
        user_count = data.get("userCount", 0)

        txBox = slide.shapes.add_textbox(Inches(0.7), Inches(4.76), Inches(3), Inches(0.3))
        p = txBox.text_frame.paragraphs[0]
        p.text = f"N = {user_count}"
        p.font.size = Pt(8)
        p.font.color.rgb = THEME_COLORS['text']

        # --- User Icon ---
        icon_path = data.get('userLogo')
        if icon_path:
            img_stream = _image_stream(icon_path)
            if img_stream:
                slide.shapes.add_picture(img_stream, Inches(0.5), Inches(4.76), width=Inches(0.25))

        # Category axis with auto font size - USE SAME SIZE FOR DATA LABELS
        if num_cats > 40:
            font_size = 5
        elif num_cats > 30:
            font_size = 6
        elif num_cats > 20:
            font_size = 7
        else:
            font_size = 8

        category_axis = chart.category_axis
        category_axis.tick_labels.font.size = Pt(font_size)
        category_axis.tick_labels.orientation = -45 if longest_label > 12 else 0

        plot = chart.plots[0]
        plot.has_data_labels = True
        data_labels = plot.data_labels
        data_labels.position = XL_DATA_LABEL_POSITION.CENTER
        data_labels.font.size = Pt(font_size)

        for series in plot.series:
            for idx, point in enumerate(series.points):
                value = series.values[idx]
                data_label = point.data_label
                if abs(value) < 0.0001:
                    data_label.text_frame.text = "0%"
                else:
                    data_label.text_frame.text = f"{value:.0%}"
                    data_label.font.size = Pt(font_size)
                    data_label.font.color.rgb = RGBColor(0, 0, 0)
                    data_label.text_frame.word_wrap = True

        # Hide value axis
        value_axis = chart.value_axis
        value_axis.tick_labels.font.size = Pt(font_size)
        value_axis.visible = False
        value_axis.has_major_gridlines = False
        value_axis.has_minor_gridlines = False

        static_colors = ["#08AE50", "#FFD735", "#E34D43", "#e02222"]
        for i, series in enumerate(chart.series):
            rgb_tuple = _hex_to_rgb(static_colors[i % len(static_colors)].lstrip("#"))
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = RGBColor(*rgb_tuple)

            for point in series.points:
                if point.data_label: 
                    point.data_label.font.size = Pt(font_size)

        return slide
    except Exception as e:
        print(f"Error in function '_create_stacked_bar_chart_slide': {e}")
        raise

def _create_hra_summary_slide(prs, data):
    try:
        MAX_ROWS_PER_SLIDE = 14
        slides_created = []

        table_data = data.get('rows', [])
        headers = data.get('headers', [])

        def chunk_rows(rows, size):
            for i in range(0, len(rows), size):
                yield rows[i:i + size]

        for chunk_idx, rows_chunk in enumerate(chunk_rows(table_data, MAX_ROWS_PER_SLIDE)):
            # --- NEW SLIDE ---
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slides_created.append(slide)
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor(241, 252, 248)

            # --- TITLE ---
            title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.5))
            p = title_box.text_frame.paragraphs[0]
            title_text = data.get('title', 'HRA Summary Results')
            if chunk_idx > 0:
                title_text += f" "
            p.text = title_text
            p.font.size = Pt(20)
            p.font.bold = True
            p.font.color.rgb = THEME_COLORS['accent_gray']

            # --- TABLE ---
            if headers and rows_chunk:
                rows_count = len(rows_chunk) + 1
                cols_count = len(headers)
                table_shape = slide.shapes.add_table(
                    rows_count, cols_count, Inches(0.5), Inches(0.8), Inches(9.0), Inches(3.5)
                )
                table = table_shape.table

                # --- FIXED ROW HEIGHTS ---
                table.rows[0].height = Inches(0.27)  # Header height
                for row_idx in range(1, rows_count):
                    table.rows[row_idx].height = Inches(0.26)  # Data row height

                # Header row
                for col_idx, header in enumerate(headers):
                    cell = table.cell(0, col_idx)
                    cell.text = str(header)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(0x2D, 0x3C, 0x3F)
                    p = cell.text_frame.paragraphs[0]
                    p.clear()
                    run = p.add_run()
                    run.text = str(header)
                    run.font.bold = True
                    run.font.size = Pt(8)
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    p.alignment = PP_ALIGN.CENTER

                # Data rows
                for row_idx, row_data in enumerate(rows_chunk):
                    rgb_tuple = _hex_to_rgb("F1FCF8")
                    bg_color = RGBColor(*rgb_tuple)
                    for col_idx, value in enumerate(row_data):
                        cell = table.cell(row_idx + 1, col_idx)
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = bg_color
                        p = cell.text_frame.paragraphs[0]
                        p.clear()
                        run = p.add_run()
                        run.text = str(value)
                        run.font.size = Pt(8)
                        run.font.color.rgb = THEME_COLORS['text']
                        p.alignment = PP_ALIGN.CENTER if col_idx > 0 else PP_ALIGN.LEFT

            # --- FOOTER ---
            user_count = data.get('userCount')
            if user_count:
                footer_text = f"N = {user_count}"
                footer_box = slide.shapes.add_textbox(Inches(0.7), Inches(4.76), Inches(3), Inches(0.3))
                p = footer_box.text_frame.paragraphs[0]
                p.text = footer_text
                p.font.size = Pt(8)
                p.font.color.rgb = THEME_COLORS['text']

            # --- ICON ---
            icon_path = data.get('iconPath')
            if icon_path:
                img_stream = _image_stream(icon_path)
                if img_stream:
                    slide.shapes.add_picture(img_stream, Inches(0.1), Inches(4.65), Inches(0.4), Inches(0.4))

            # --- User Icon ---
            icon_path = data.get('userLogo')
            if icon_path:
                img_stream = _image_stream(icon_path)
                if img_stream:
                    slide.shapes.add_picture(img_stream, Inches(0.5), Inches(4.76), width=Inches(0.25))

        return slides_created
    except Exception as e:
        print(f"Error in function '_create_hra_summary_slide': {e}")
        raise
    

def _create_hra_results_chart_slide(prs, data):
    try:
        slide = _add_standard_header_slide(prs, data.get("title", "HRA Results"), "")
        chart_data = CategoryChartData()
        categories = data.get("categories", [])
        chart_data.categories = categories

        def clean_data(series_data):
            cleaned = []
            for value in series_data:
                try:
                    cleaned.append(float(str(value).strip().replace('%', '')) / 100.0)
                except (ValueError, TypeError):
                    cleaned.append(0.0)
            return cleaned

        chart_data.add_series("Low Risk", clean_data(data.get("low", [])))
        chart_data.add_series("Medium Risk", clean_data(data.get("medium", [])))
        moderate_count = data.get("moderate", [])
        if moderate_count:
            chart_data.add_series("Moderate Risk", clean_data(data.get("moderate", [])))
        
        chart_data.add_series("High Risk", clean_data(data.get("high", [])))
        very_high_count = data.get("very high", [])
        if very_high_count:
            chart_data.add_series("Very High Risk", clean_data(data.get("very high", [])))

        slide_width = prs.slide_width
        slide_height = prs.slide_height

        longest_label = max((len(str(c)) for c in categories), default=10)
        est_label_width = Inches(0.12 * longest_label)  # ~0.12 inch per char at 8pt

        left = Inches(0.5) + est_label_width
        top = Inches(0.8)
        width = slide_width - left - Inches(0.5)  

        # Dynamic height based on number of categories
        num_cats = len(categories)
        base_height = Inches(4.6)
        min_height = Inches(2.5)
        max_height = slide_height - Inches(2.0)  # leave space for header/footer

        if num_cats > 0:
            height = base_height * (num_cats / 10)  # scale relative to 10 categories
        else:
            height = base_height

        height = max(min_height, min(max_height, height))

        chart_type = XL_CHART_TYPE.BAR_STACKED_100
        chart_shape = slide.shapes.add_chart(chart_type, left, top, width, height, chart_data)
        chart = chart_shape.chart

        # Legend styling (global only, no per-entry support in python-pptx)
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.font.size = Pt(6)

        # Category axis with auto font size - USE SAME SIZE FOR DATA LABELS
        if num_cats > 40:
            font_size = 5
        elif num_cats > 30:
            font_size = 6
        elif num_cats > 20:
            font_size = 7
        else:
            font_size = 8

        category_axis = chart.category_axis
        category_axis.tick_labels.font.size = Pt(font_size)
        category_axis.tick_labels.orientation = -45 if longest_label > 12 else 0

        plot = chart.plots[0]
        plot.has_data_labels = True
        data_labels = plot.data_labels
        data_labels.position = XL_DATA_LABEL_POSITION.CENTER
        data_labels.font.size = Pt(font_size)

        for series in plot.series:
            for idx, point in enumerate(series.points):
                # value = series.values[idx]
                value = series.values[idx] if idx < len(series.values) else 0
                data_label = point.data_label
                if abs(value) < 0.0001:
                    data_label.text_frame.text = "0%"
                else:
                    data_label.text_frame.text = f"{value:.0%}"
                    data_label.font.size = Pt(font_size)
                    data_label.font.color.rgb = RGBColor(0, 0, 0)
                    data_label.text_frame.word_wrap = True

        # Hide value axis
        value_axis = chart.value_axis
        value_axis.tick_labels.font.size = Pt(font_size)
        value_axis.visible = False
        value_axis.has_major_gridlines = False
        value_axis.has_minor_gridlines = False

        static_colors = ["#08AE50", "#FFD735", "#FE0000"]
        if moderate_count:
            static_colors.insert(1, "#FFA500")
        if very_high_count:
            static_colors.append("#690202")
        for i, series in enumerate(chart.series):
            rgb_tuple = _hex_to_rgb(static_colors[i % len(static_colors)].lstrip("#"))
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = RGBColor(*rgb_tuple)

            for point in series.points:
                if point.data_label: 
                    point.data_label.font.size = Pt(font_size)

        # --- FOOTER ---
        user_count = data.get('userCount')
        if user_count:
            footer_text = f"N = {user_count}"
            footer_box = slide.shapes.add_textbox(Inches(0.9), Inches(4), Inches(3), Inches(0.3))
            p = footer_box.text_frame.paragraphs[0]
            p.text = footer_text
            p.font.size = Pt(8)
            p.font.color.rgb = THEME_COLORS['text']
        # --- User Icon ---
        icon_path = data.get('userLogo')
        if icon_path:
            img_stream = _image_stream(icon_path)
            if img_stream:
                slide.shapes.add_picture(img_stream, Inches(0.7), Inches(4), width=Inches(0.25))

        return slide
    except Exception as e:
        print(f"Error in function '_create_hra_results_chart_slide': {e}")
        raise

def _create_risk_factors_slide(prs, data):
    try:
        # --- 1. Create Slide and Set Background (similar to _create_table_slide) ---
        title = data.get('title', 'Number Of Users With High Risk Factors')
        slide = _add_standard_header_slide(prs, title, subtitle=None)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(240, 248, 247)

        # --- 2. Prepare Data for Both Tables ---
        # Left Table Data
        risk_factors_list = data.get('riskFactorsList', [])
        risk_factors_table_data = [['Risk Factors']] + [[factor] for factor in risk_factors_list]
        
        # Right Table Data
        dist_data = data.get('riskDistribution', {})
        sorted_years = sorted(dist_data.keys(), reverse=True)
        headers = ["Number of High Risk Factors"] + [f"{year} (#)" for year in sorted_years]
        row_labels = ["0", "1", "2", "3", "4 +"]
    
        table_rows_right = [headers]
        for i, label in enumerate(row_labels):
            row = [label]
            for year in sorted_years:
                year_data = dist_data.get(year, [])
                value = year_data[i] if i < len(year_data) else 0
                row.append(value)
            table_rows_right.append(row)

        # --- 3. Create Left Table ("Risk Factors") ---
        rows_left = len(risk_factors_table_data)
        cols_left = 1
        left_pos, top_pos = Inches(0.4), Inches(1.3)
        width_left, height_left = Inches(1.8), Inches(rows_left * 0.3)

        # Create the table itself
        table_left = slide.shapes.add_table(rows_left, cols_left, left_pos, top_pos, width_left, height_left).table
        table_left.columns[0].width = width_left

        # Populate and style the left table
        for r, row_data in enumerate(risk_factors_table_data):
            cell = table_left.cell(r, 0)
            cell.text = row_data[0]
            cell.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
            
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            p.font.size = Pt(8)

            if r == 0:  # Header row styling
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(30, 53, 58) 
                p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)   
                p.font.bold = True
            else:  # Data row styling
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(240, 248, 247) 
                p.font.color.rgb = RGBColor(0x0a, 0x41, 0x51)   
                p.font.bold = False
                set_cell_border(cell, hex_color="#AAAAAA", border_type="dotted", sides=("bottom",))


        # --- 4. Create Right Table ("Distribution") ---
        rows_right = len(table_rows_right)
        cols_right = len(headers)
        left_pos, top_pos = Inches(2.8), Inches(1.3)
        width_right, height_right = Inches(5.7), Inches(rows_right * 0.41)

        table_right = slide.shapes.add_table(rows_right, cols_right, left_pos, top_pos, width_right, height_right).table
        
        # Set column widths
        table_right.columns[0].width = Inches(1.3)
        if cols_right > 1:
            col_width = Inches(5.7 / (cols_right - 1))
            for i in range(1, cols_right):
                table_right.columns[i].width = col_width
        
        # Populate and style the right table
        for r, row_data in enumerate(table_rows_right):
            for c, cell_data in enumerate(row_data):
                cell = table_right.cell(r, c)
                cell.text = str(cell_data)
                cell.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
                p = cell.text_frame.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER

                if r == 0:  # Header row styling
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(30, 53, 58)
                    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    p.font.bold = True
                    p.font.size = Pt(8)
                else:  # Data row styling
                    cell.fill.background()
                    p.font.color.rgb = RGBColor(0x0a, 0x41, 0x51)
                    p.font.size = Pt(8)
                    set_cell_border(cell, hex_color="#AAAAAA", border_type="dotted", sides=("bottom",))

        # --- 5. Add Footer (similar to _create_table_slide) ---
        user_count = data.get('userCount', 0)

        txBox = slide.shapes.add_textbox(Inches(0.7), Inches(4.76), Inches(3), Inches(0.3))
        p = txBox.text_frame.paragraphs[0]
        p.text = f"N = {user_count}"
        p.font.size = Pt(8)
        p.font.color.rgb = THEME_COLORS['text']
        # --- User Icon ---
        icon_path = data.get('userLogo')
        if icon_path:
            img_stream = _image_stream(icon_path)
            if img_stream:
                slide.shapes.add_picture(img_stream, Inches(0.5), Inches(4.76), width=Inches(0.25))

        return slide
    except Exception as e:
        print(f"Error in function '_create_risk_factors_slide': {e}")
        raise

def _create_questions_slide(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # --- Background Color ---
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(30, 53, 58)

    # --- BACKGROUND IMAGE ---
    bg_path = data.get('backgroundImagePath')
    if bg_path:
        img_stream = _image_stream(bg_path)
        if img_stream is not None:
            slide.shapes.add_picture(img_stream, Inches(0), Inches(2.63), Inches(10), Inches(3))

    # --- Title "Questions & Comments?" ---
    # 1. Create Shape
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.7), Inches(5), Inches(0.8))
    
    # 2. Access Paragraph
    p_title = title_box.text_frame.paragraphs[0]
    
    # 3. CRITICAL FIX: Add a Run explicitly
    run_title = p_title.add_run()
    
    # 4. Apply styles to the RUN, not the paragraph
    run_title.text = data.get("title", "Questions & Comments?")
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(255, 255, 255)

    # --- Contact Info ---
    default_contact = """
    1700 Post Oak Boulevard,
    Suite 600 Houston, TX 77056
    1-877-378-8880
    
    info@zomohealth.com
    zomohealth.com
    """
    contact_text = data.get("contactInfo", textwrap.dedent(default_contact).strip())

    contact_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(5), Inches(2.5))
    contact_box.text_frame.word_wrap = True
    
    p_contact = contact_box.text_frame.paragraphs[0]
    
    # CRITICAL FIX: Use a Run here as well
    run_contact = p_contact.add_run()
    run_contact.text = contact_text
    run_contact.font.size = Pt(10)
    run_contact.font.color.rgb = RGBColor(255, 255, 255)

    # --- Site Logo ---
    logo_path = data.get("siteLogo")
    if logo_path:
        img_stream = _image_stream(logo_path)
        if img_stream is not None:
            slide.shapes.add_picture(img_stream, Inches(7.5), Inches(0.5), Inches(1.82), Inches(0.20))

    return slide


# --- MAIN PRESENTATION GENERATOR ---
def generate_presentation(report_data, output_path):
    prs = pptx.Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)

    slide_generators = {
        "title": _create_title_slide,
        "sectionBreak": _create_section_break_slide,
        "engagementSummary": _create_table_slide,
        "historicalTable": _create_table_slide,
        "varianceTable": _create_table_slide,
        "multiYearAverageTable": _create_table_slide,
        "stackedBarChart": _create_stacked_bar_chart_slide,
        "riskUserTable": _create_risk_factors_slide,
        "hraSummary": _create_hra_summary_slide,
        "hraResultChart": _create_hra_results_chart_slide,
        "contact": _create_questions_slide,
    }

    total_slides_data = report_data.get('slides', [])
    total_slides = len(total_slides_data)

    slide_counter = 0  # track real slides created

    for idx, slide_data in enumerate(total_slides_data):
        slide_type = slide_data.get('slideType')
        if slide_type in slide_generators:
            generated = slide_generators[slide_type](prs, slide_data.get('data', {}))

            # Handle multiple slides (tables return list)
            slides = generated if isinstance(generated, list) else [generated]

            for s in slides:
                hide_page = (slide_type in ["title", "sectionBreak", "contact"])
                is_title_slide = (slide_type == "title")
                if slide_type == "sectionBreak":
                    footer_logo_path = report_data.get('footerLogo')
                else:
                    footer_logo_path = report_data.get('footerLogoDark')

                _add_footer(prs, s, slide_counter, total_slides, footer_logo_path, hide_page, is_title_slide)
                slide_counter += 1

        else:
            print(f"Warning: No generator found for slide type '{slide_type}'", file=sys.stderr)

    prs.save(output_path)


# --- SCRIPT EXECUTION ---
if __name__ == "__main__":
    try:
        if len(sys.argv) < 2:
            raise ValueError("Output file path argument is missing.")
        output_file_path = sys.argv[1]
        input_data_string = sys.stdin.read()
        if not input_data_string:
            raise ValueError("No data received from stdin.")
        dynamic_report_data = json.loads(input_data_string)
        generate_presentation(dynamic_report_data, output_file_path)
    except Exception as e:
        print(f"Python Error: {e}", file=sys.stderr)
        sys.exit(1)
