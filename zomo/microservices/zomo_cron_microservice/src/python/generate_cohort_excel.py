import openpyxl
import json
import tempfile
import os
import sys
import boto3
import traceback
from openpyxl.cell.cell import MergedCell

def error_response(msg, err):
    print("ERROR:", msg)
    print("TYPE:", type(err).__name__)
    print("DETAIL:", str(err))
    print("TRACEBACK:")
    print(traceback.format_exc())
    sys.exit(1)

def get_s3_client(region):
    local = os.environ.get("LOCAL")
    if local == "true":
        key = os.environ.get("AWS_ACCESS_KEY_PROD")
        secret = os.environ.get("AWS_SECRET_ACCESS_KEY_PROD")
        return boto3.client(
            "s3",
            aws_access_key_id=key,
            aws_secret_access_key=secret,
            region_name=region
        )
    return boto3.client("s3", region_name=region)

def main():
    try:
        if len(sys.argv) < 2:
            raise Exception("Missing argument: JSON path (S3 key) not provided")
        S3_JSON_KEY = sys.argv[1]
        AWS_REGION = os.environ.get("AWS_REGION", "us-east-1")
        S3_BUCKET = os.environ.get("AWS_BUCKET_PRIVATE_PROD")
        S3_TEMPLATE_KEY = os.environ.get("TEMPLATE_PATH")
        if not all([S3_BUCKET, S3_TEMPLATE_KEY]):
            raise Exception("Missing required AWS environment variables")
        s3 = get_s3_client(AWS_REGION)
        with tempfile.NamedTemporaryFile(delete=False, mode="w+b") as temp_json:
            s3.download_fileobj(S3_BUCKET, S3_JSON_KEY, temp_json)
            temp_json_path = temp_json.name
        with open(temp_json_path, "r", encoding="utf-8") as f:
            data = json.load(f)
        headers = data.get("clm_name_arr", [])
        rows = data.get("clm_data_arr", [])
        sheet_titles = data.get("sheet_title", [])
        xlsx_key = data.get("filename")
        if not xlsx_key:
            raise Exception("JSON missing 'filename'")
        with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as temp_template:
            s3.download_fileobj(S3_BUCKET, S3_TEMPLATE_KEY, temp_template)
            template_path = temp_template.name
        wb = openpyxl.load_workbook(template_path)
        for i, title in enumerate(sheet_titles):
            if i >= len(wb.sheetnames):
                wb.create_sheet(title)
            ws = wb.worksheets[i]
            ws.title = title
            if i < len(headers):
                for col_idx, header in enumerate(headers[i]):
                    ws.cell(row=1, column=col_idx + 1, value=header)
            if i < len(rows):
                for row_idx, rowdata in enumerate(rows[i]):
                    for col_idx, value in enumerate(rowdata):
                        cell = ws.cell(row=row_idx + 2, column=col_idx + 1)
                        if not isinstance(cell, MergedCell):
                            cell.value = value
        with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as temp_xlsx:
            wb.save(temp_xlsx.name)
            temp_xlsx_path = temp_xlsx.name
        with open(temp_xlsx_path, "rb") as f:
            s3.upload_fileobj(f, S3_BUCKET, xlsx_key)
        print(f"SUCCESS: File uploaded to s3://{S3_BUCKET}/{xlsx_key}")
    except Exception as e:
        error_response("Unexpected failure", e)
    finally:
        try:
            if 'temp_json_path' in locals() and os.path.exists(temp_json_path):
                os.remove(temp_json_path)
            if 'template_path' in locals() and os.path.exists(template_path):
                os.remove(template_path)
            if 'temp_xlsx_path' in locals() and os.path.exists(temp_xlsx_path):
                os.remove(temp_xlsx_path)
        except:
            pass

if __name__ == "__main__":
    main()
